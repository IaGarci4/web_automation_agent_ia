"""
Genera el archivo de casos de TRN-239 listo para importar en QMetry.

Produce **.xlsx y .csv** con las 21 columnas en el orden que espera QMetry
(metadatos en la primera fila del caso, un paso por fila).

    python tools/trn239_qmetry_xlsx.py
    python tools/trn239_qmetry_xlsx.py --salida C:\\ruta\\donde\\quieras

## Por qué pocos casos

La evidencia real de cada caso es **el mismo reporte HTML**, que desglosa paso
a paso lo que aquí se resume. Documentar en QMetry lo que el reporte ya cuenta
mejor crea dos versiones de la verdad, y la de QMetry es la que envejece. Los
pasos de aquí describen **qué se ejecuta**, no cada aserción.

Cinco casos, cada uno agrupa lo que se ejecuta y se lee junto:

    CP01  el servicio responde
    CP02  el camino feliz: alta y cancelación, JSON y SOAP
    CP03  lo que la API debe rechazar (negativos controlados)
    CP04  lo que decide la migración: los datos guardados y su coherencia
    CP05  catálogo de errores: cada Errorcode se dispara a propósito
"""

import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

TICKET = "TRN-239"
COMPONENTE = "URANO"          # lo nombra explícitamente la doc del ticket
AMBIENTE = "TEST"
TIPO = "Functional"
ESTADO = "IN PROGRESS"
VERSION = ""
DEPLOY = "29/SEP/2026"        # OR-960
AI_GENERATED = "SI"
UUID = "6410df3d7222b08f3e70b344"        # Ignacio Antonio Garcia Espinoza

REPORTE = r"reports\evidence\TRN-239\reporte_TRN-239.html"
EVIDENCIA = (f"Evidencia: el reporte consolidado de la corrida, {REPORTE}. "
             f"Incluye el par petición/respuesta de cada llamada y las "
             f"lecturas de base de datos.")

COLUMNAS = [
    "Summary", "Description", "Precondition", "Status", "Priority",
    "Assignee", "Reporter", "Estimated Time", "Labels", "Components",
    "Step Summary", "Test Data", "Expected Result", "Version", "Ambiente",
    "Tester", "Note", "TestCase Type", "AI Generated", "Story Linkages",
    "Deploy",
]

PRE = ("La API Lunex responde en "
       "https://zeus-services.maxilabs.net/api/v1/lunex/ (VPN conectada). "
       "Credenciales de TEST confirmadas. Acceso de lectura a "
       "lunex.TransferLN para los pasos de integridad.")

CMD = "pytest src/tests/etiquetas/TRN_239 -v -s"

# (id, resumen, descripción, prioridad, tiempo, pasos[(paso, datos, esperado)])
CASOS = [
    (
        "CP01", "El servicio responde y reporta el estado de su base de datos",
        "Healthcheck de la API migrada. Distingue el estado del servicio del "
        "de su base de datos: si la API contesta pero no ve la base, una alta "
        "podría responder Errorcode 0 sin guardar nada. Subticket TRN-569.",
        "Medium", "00:03:00",
        [
            ("Invocar el healthcheck con la VPN conectada.",
             CMD + " -k CP01",
             "HTTP 200, status = UP y database = UP (en details.database). "
             + EVIDENCIA),
        ],
    ),
    (
        "CP02", "Alta y cancelación de transacción, en JSON y en SOAP",
        "El camino feliz completo. Registra transacciones por los dos "
        "formatos que soporta la API (el legacy Urano hablaba SOAP, así que "
        "ese camino debe seguir vivo) y las cancela. Se comprueba que ambos "
        "formatos den el mismo resultado.",
        "Highest", "00:15:00",
        [
            ("Registrar un alta en JSON y otra en SOAP con la misma "
             "estructura, y verificar la paridad.",
             CMD + " -k CP02",
             "HTTP 200, Errorcode = 0, Error_text = Success en ambos "
             "formatos. La respuesta SOAP no trae Fault."),
            ("Cancelar la transacción registrada (Status = Cancelled, según "
             "el requerimiento).",
             CMD + " -k CP03",
             "HTTP 200 y Errorcode = 0. En lunex.TransferLN queda cancelada: "
             "DateOfCancel con fecha e IdStatus = 22 (LNStatus VOID). "
             + EVIDENCIA),
        ],
    ),
    (
        "CP03", "La API rechaza lo que debe, de forma controlada",
        "Los negativos: firma MD5 inválida, campos obligatorios ausentes, "
        "importe negativo y cancelación de una transacción inexistente. El "
        "criterio es que rechace con un Errorcode del catálogo y sin exponer "
        "trazas internas. Subtickets TRN-293, TRN-294, TRN-300.",
        "High", "00:10:00",
        [
            ("Enviar peticiones inválidas, cada una rompiendo un solo campo "
             "sobre un cuerpo por lo demás válido.",
             CMD + " -k \"CP04 or CP05\"",
             "Se rechazan con un Errorcode del catálogo: 4 (Session is "
             "invalid), 5 (Validation error) y 12 (Transaction does not "
             "exists). Ninguna respuesta contiene traza ni stack. "
             + EVIDENCIA),
        ],
    ),
    (
        "CP04", "Lo enviado es lo guardado, coherente con producción, sin "
                "duplicar",
        "El caso que justifica el ticket: comprueba que la base quedó bien. "
        "Compara lo enviado contra lo persistido, valida la coherencia contra "
        "las reglas de producción (Commission=D1 en ITU, Commission=Agent+"
        "Corp salvo DTU, AmountInMN=Amount×ExRate) y la idempotencia (dos "
        "filas con el mismo TransactionID serían una recarga cobrada dos "
        "veces).",
        "Highest", "00:20:00",
        [
            ("Registrar un alta, comparar campo por campo contra "
             "lunex.TransferLN y validar las reglas de producción.",
             CMD + " -k \"CP06 or CP07 or CP08\"",
             "La fila existe, los campos coinciden y las reglas duras de "
             "producción se cumplen. Un campo que no cuadra sale como FALLA "
             "con el valor exacto."),
            ("Reenviar el mismo TransactionID y contar las filas.",
             "reenvío del alta con el mismo TransactionID",
             "La API responde Errorcode 2 (Transaction already exists) y en "
             "la base sigue habiendo una sola fila. " + EVIDENCIA),
        ],
    ),
    (
        "CP05", "Catálogo de errores: cada Errorcode se dispara a propósito",
        "Regresión negativa. Se manda basura a propósito para garantizar que "
        "las validaciones (firma, esquema, unicidad, estados, agente) siguen "
        "disparando tras el cambio de desarrollo. Pasa cuando la API RECHAZA; "
        "si acepta algo que debía rechazar, se anota con los datos para "
        "reproducirlo.",
        "High", "00:12:00",
        [
            ("Enviar una petición inválida por cada Errorcode del "
             "requerimiento (1,2,3,4,5,6,7,8,9,10,12,15), request "
             "independiente por cada uno.",
             CMD + " -k CP09",
             "Cada petición se rechaza con un Errorcode documentado coherente "
             "con lo que se rompió. Los casos donde la API acepta lo inválido "
             "quedan marcados con los datos para reproducirlos. " + EVIDENCIA),
        ],
    ),
]


def filas() -> list:
    """Todas las filas: metadatos solo en la primera de cada caso."""
    salida = [COLUMNAS]
    for cid, resumen, descripcion, prioridad, tiempo, pasos in CASOS:
        for i, (paso, datos, esperado) in enumerate(pasos):
            if i == 0:
                salida.append([
                    f"{TICKET} {cid} {resumen}", descripcion, PRE, ESTADO,
                    prioridad, UUID, UUID, tiempo, TICKET, COMPONENTE,
                    paso, datos, esperado, VERSION, AMBIENTE, UUID, "",
                    TIPO, AI_GENERATED, TICKET, DEPLOY,
                ])
            else:
                # Filas de paso: solo las tres columnas del paso. Rellenar los
                # metadatos aquí haría que QMetry creara un caso por fila.
                salida.append(["", "", "", "", "", "", "", "", "", "",
                               paso, datos, esperado,
                               "", "", "", "", "", "", "", ""])
    return salida


def escribir_csv(destino: Path, datos: list) -> None:
    with open(destino, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(datos)


def escribir_xlsx(destino: Path, datos: list) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"
    for fila in datos:
        ws.append(fila)

    encabezado = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1F4E78")
    for celda in ws[1]:
        celda.font = encabezado
        celda.fill = fondo
        celda.alignment = Alignment(vertical="center")

    anchos = {"A": 58, "B": 70, "C": 55, "D": 14, "E": 11, "F": 26, "G": 26,
              "H": 15, "I": 12, "J": 14, "K": 58, "L": 42, "M": 78, "N": 10,
              "O": 12, "P": 26, "Q": 22, "R": 16, "S": 13, "T": 15, "U": 14}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(destino)
    return True


def main() -> int:
    ap = argparse.ArgumentParser(description="Casos de TRN-239 para QMetry")
    ap.add_argument("--salida",
                    default=str(Path(__file__).resolve().parents[1]
                                / "reports" / "evidence" / "TRN-239"))
    args = ap.parse_args()

    destino = Path(args.salida)
    destino.mkdir(parents=True, exist_ok=True)
    datos = filas()

    ruta_csv = destino / f"{TICKET}-QMetry.csv"
    escribir_csv(ruta_csv, datos)
    ruta_xlsx = destino / f"{TICKET}-QMetry.xlsx"
    hay_xlsx = escribir_xlsx(ruta_xlsx, datos)

    print()
    print("═" * 66)
    print(f"  {TICKET} · casos para QMetry")
    print("═" * 66)
    print(f"  {len(CASOS)} casos · {sum(len(c[5]) for c in CASOS)} pasos · "
          f"{len(COLUMNAS)} columnas")
    print(f"  Componente {COMPONENTE} · Ambiente {AMBIENTE} · Deploy {DEPLOY}")
    print("  Evidencia de todos: el mismo reporte HTML de la corrida")
    print("─" * 66)
    print(f"  CSV : {ruta_csv}")
    if hay_xlsx:
        print(f"  XLSX: {ruta_xlsx}")
    else:
        print("  XLSX: NO generado — falta openpyxl.")
        print("        pip install openpyxl   y vuelve a correr esto.")
        print("        Mientras tanto, el CSV se importa igual en QMetry.")
    print("═" * 66)
    print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
